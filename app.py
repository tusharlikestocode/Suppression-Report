import streamlit as st
import pandas as pd
from datetime import datetime
from pptx import Presentation
from pptx.util import Inches
import tempfile

st.title("📊 CSV to Insights")

uploaded_file = st.file_uploader("Upload CSV", type=["csv"])

if uploaded_file:

    df = pd.read_csv(uploaded_file)

    # --- CLEANING ---
    df['NEXT_BILLING_DATE'] = pd.to_datetime(df['NEXT_BILLING_DATE'], errors='coerce')
    today = datetime.today()

    # Filter dates
    df = df[df['NEXT_BILLING_DATE'].dt.year != 2025]
    df = df[
        (df['NEXT_BILLING_DATE'].dt.year < today.year) |
        (
            (df['NEXT_BILLING_DATE'].dt.year == today.year) &
            (df['NEXT_BILLING_DATE'].dt.month <= today.month)
        )
    ]

    # Deduplicate
    df = df.sort_values(by='NEXT_BILLING_DATE', ascending=False)
    df = df.drop_duplicates(subset=['ACCOUNT_NO'], keep='first')

    # Remove dummy
    df = df[~df['NOTE_TEXT'].str.lower().str.contains('dummy', na=False)]

    # Categorize
    def categorize(note):
        note = str(note).lower()
        if "space out" in note:
            return "Space Out"
        elif "adj" in note:
            return "Biz planned suppression (pending Adjustment)"
        elif "pending order" in note:
            return "Biz planned suppression (pending order)"
        else:
            return "Biz planned suppression"
    df['category'] = df['NOTE_TEXT'].apply(categorize)

    # Month
    df['month'] = df['NEXT_BILLING_DATE'].dt.strftime('%b %Y')

    # --- INSIGHTS ---
    grand_total = len(df)

    month_total = df.groupby('month').size().reset_index(name='total_count')
    month_total = month_total.sort_values(
    by='month',
    key=lambda x: pd.to_datetime(x, format='%b %Y')
)
    category_total = df.groupby('category').size().reset_index(name='total_count')

    # Grouped category
    df['category_grouped'] = df['category'].apply(
        lambda x: 'Space Out' if x == 'Space Out' else 'Business Suppression'
    )

    grouped_summary = df.groupby(['month', 'category_grouped']).size().reset_index(name='count')
    grouped_summary = grouped_summary.sort_values(
    by='month',
    key=lambda x: pd.to_datetime(x, format='%b %Y')
)

    # --- DISPLAY ---
    st.subheader("📌Total Suppression Count")
    st.write(grand_total)

    st.subheader("📅 Month-wise Suppression Count")
    st.dataframe(month_total)

    st.subheader("📊 Category Total")
    st.dataframe(category_total)

    st.subheader("📊 Space Out vs Business Suppression")
    st.dataframe(grouped_summary)
    # --- PPT GENERATION ---
    if st.button("Generate PPT"):

        prs = Presentation()

        # Title slide
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "CSV Insights Report"

        # Summary slide
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Summary"
        content = slide.placeholders[1]

        text = f"Total Records: {grand_total}\n\nCategory Breakdown:\n"
        for _, row in category_total.iterrows():
            text += f"{row['category']}: {row['total_count']}\n"

        content.text = text

        # Save temp file
        tmp = tempfile.NamedTemporaryFile(delete=False, suffix=".pptx")
        prs.save(tmp.name)

        with open(tmp.name, "rb") as f:
            st.download_button(
                "⬇ Download PPT",
                f,
                file_name="report.pptx"
            )